"""Generate only synthetic parser acceptance files in a new directory.

Run with a development Python containing reportlab, python-docx, openpyxl and
python-pptx: ``python create_file_fixtures.py /tmp/new-file-fixtures``. These
fixture-generation dependencies are not installed in the API or task runtime.
"""

import argparse
from pathlib import Path

MARKER = "FILE_RUNTIME_271828"


def create_fixtures(directory: Path) -> None:
    """Create Chinese Office/PDF, partial PDF and encrypted PDF without overwriting files."""

    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.lib.pdfencrypt import StandardEncryption
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen.canvas import Canvas

    directory.mkdir(mode=0o700, parents=False, exist_ok=False)
    text = f"中文文件读取测试 {MARKER}"
    word = Document()
    word.add_paragraph(text)
    word.add_table(rows=1, cols=2).cell(0, 0).text = "表格内容"
    word.save(directory / "sample.docx")
    workbook = Workbook()
    workbook.active.title = "数据"
    workbook.active.append([text, 42])
    workbook.save(directory / "sample.xlsx")
    workbook.close()
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = text
    slide.notes_slide.notes_text_frame.text = "备注测试"
    slides.save(directory / "sample.pptx")
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    for name, encrypted in (("sample.pdf", False), ("encrypted.pdf", True)):
        canvas = Canvas(str(directory / name), pagesize=(600, 300),
                        encrypt=StandardEncryption("fixture-password") if encrypted else None)
        canvas.setFont("STSong-Light", 16)
        canvas.drawString(30, 220, text)
        canvas.showPage()
        canvas.showPage()  # Deliberately empty: must be recorded as uncovered page 2.
        canvas.setFont("STSong-Light", 16)
        canvas.drawString(30, 220, "第三页文本")
        canvas.showPage()
        canvas.save()


def main() -> None:
    """Accept only the destination for a new set of synthetic development fixtures."""

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("directory", type=Path)
    create_fixtures(parser.parse_args().directory)


if __name__ == "__main__":
    main()
