"""Dispatch fixed Office/PDF parsing inside the offline disposable file-runtime container.

Input is one bounded office_parse.v1 or file_parse.v1 JSON line, containing only MIME
and base64 bytes. Output is one bounded JSON result; no paths, commands or URLs
are accepted. Run through the administrator-owned executor, never on API hosts.
"""

import base64
import hashlib
import importlib.metadata
import io
import json
import re
import resource
import stat
import sys
import zipfile
from datetime import date, datetime, time
from pathlib import PurePosixPath
from xml.etree import ElementTree

MAX_SOURCE_BYTES = 20 * 1024 * 1024
# Include base64 expansion and fixed JSON protocol fields, not just source bytes.
MAX_REQUEST_BYTES = ((MAX_SOURCE_BYTES + 2) // 3) * 4 + 4096
MAX_RESULT_BYTES = 4 * 1024 * 1024
MAX_BLOCKS = 10000
OFFICE_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "word/document.xml",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xl/workbook.xml",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": (
        "ppt/presentation.xml"
    ),
}


def validate_package(content: bytes, media_type: str) -> None:
    """Reject mismatched, encrypted, excessive, unsafe or active OOXML ZIP packages."""

    if media_type not in OFFICE_TYPES or not 0 < len(content) <= MAX_SOURCE_BYTES:
        raise ValueError("Unsupported Office format or source size")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        names = [entry.filename for entry in entries]
        if len(entries) > 2048 or len(set(names)) != len(names):
            raise ValueError("Office package has excessive or duplicate entries")
        if sum(entry.file_size for entry in entries) > 20 * 1024 * 1024:
            raise ValueError("Office package expands beyond the supported size")
        if OFFICE_TYPES[media_type] not in names or "[Content_Types].xml" not in names:
            raise ValueError("Office package does not match its declared format")
        for entry in entries:
            path = PurePosixPath(entry.filename)
            mode = entry.external_attr >> 16
            if (
                path.is_absolute()
                or ".." in path.parts
                or "\\" in entry.filename
                or entry.flag_bits & 1
                or stat.S_ISLNK(mode)
                or entry.file_size > 200 * max(entry.compress_size, 1)
                # Libraries retain printer settings as opaque bytes; they do not execute them.
                or (
                    entry.filename.lower().endswith(".bin")
                    and not re.fullmatch(
                        r"(?:word|ppt)/printerSettings/printerSettings[0-9]+\.bin",
                        entry.filename,
                    )
                )
                or "embeddings" in path.parts
            ):
                raise ValueError("Office package contains unsafe or unsupported entries")
            if entry.filename.endswith((".xml", ".rels")):
                xml = archive.read(entry)
                # OOXML XML is textual; remove NULs to detect declarations in UTF-16 too.
                normalized = xml.replace(b"\x00", b"").upper()
                if b"<!DOCTYPE" in normalized or b"<!ENTITY" in normalized:
                    raise ValueError("Office XML declarations are not supported")
        types = ElementTree.fromstring(archive.read("[Content_Types].xml"))
        declared = {item.attrib.get("ContentType", "") for item in types}
        if any("macroenabled" in value.lower() for value in declared):
            raise ValueError("Macro-enabled Office documents are not supported")


def append_block(blocks: list[dict], text: str, **positions) -> None:
    """Append a bounded block with source-native locations, never invented text line numbers."""

    if not text.strip():
        return
    block = {"block_id": f"block:{len(blocks) + 1}", "text": text, **positions}
    if len(blocks) >= MAX_BLOCKS or len(json.dumps(block).encode()) > 128 * 1024:
        raise ValueError("Office content exceeds the block or row limit")
    blocks.append(block)


def parse_word(content: bytes) -> tuple[list[dict], None]:
    """Read body paragraphs and table rows in document order; other stories remain omitted."""

    from docx import Document
    from docx.table import Table

    document = Document(io.BytesIO(content))
    blocks = []
    paragraph_number = table_number = 0
    for element in document.iter_inner_content():
        if isinstance(element, Table):
            table_number += 1
            for row_number, row in enumerate(element.rows, 1):
                cells = [cell.text for cell in row.cells]
                append_block(
                    blocks,
                    json.dumps(cells, ensure_ascii=False),
                    kind="table_row",
                    cells=cells,
                    table_number=table_number,
                    row_number=row_number,
                )
        else:
            paragraph_number += 1
            level = None
            style_id = element.style.style_id if element.style else ""
            if style_id.startswith("Heading") and style_id[7:].isdigit():
                candidate = int(style_id[7:])
                if 1 <= candidate <= 9:
                    level = candidate
            append_block(
                blocks,
                element.text,
                paragraph_number=paragraph_number,
                heading_level=level,
            )
    return blocks, None


def cell_text(value) -> str:
    """Serialize cell values without evaluating formulas or guessing number/date semantics."""

    if value is None:
        return ""
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def parse_excel(content: bytes) -> tuple[list[dict], dict]:
    """Inventory every worksheet and extract cells without calculating formulas or charts."""

    from openpyxl import load_workbook

    original = load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(content), data_only=True, keep_links=False)
    blocks = []
    entries = []
    scanned_cells = 0
    try:
        for sheet in original:
            first_block = len(blocks)
            scanned_cells += sheet.max_row * sheet.max_column
            if sheet.max_row > 10000 or sheet.max_column > 1024 or scanned_cells > 200000:
                raise ValueError("Workbook dimensions exceed the supported reading range")
            for row in sheet.iter_rows():
                if not any(cell.value is not None for cell in row):
                    continue
                formulas = [cell.value if cell.data_type == "f" else None for cell in row]
                cells = [
                    cell_text(cached[sheet.title][cell.coordinate].value)
                    if cell.data_type == "f"
                    else cell_text(cell.value)
                    for cell in row
                ]
                append_block(
                    blocks,
                    json.dumps(cells, ensure_ascii=False),
                    kind="table_row",
                    cells=cells,
                    sheet_name=sheet.title,
                    row_number=row[0].row,
                    cell_range=f"{row[0].coordinate}:{row[-1].coordinate}",
                    cell_formulas=formulas,
                    cell_types=[cell.data_type for cell in row],
                    cell_number_formats=[cell.number_format for cell in row],
                    missing_formula_values=[
                        cell.coordinate
                        for cell in row
                        if cell.data_type == "f"
                        and cached[sheet.title][cell.coordinate].value is None
                    ],
                    merged_ranges=[
                        str(area)
                        for area in sheet.merged_cells.ranges
                        if area.min_row <= row[0].row <= area.max_row
                    ],
                    hidden=sheet.sheet_state != "visible"
                    or sheet.row_dimensions[row[0].row].hidden,
                )
            entries.append({
                "name": sheet.title, "visibility": sheet.sheet_state,
                "block_count": len(blocks) - first_block,
            })
    finally:
        original.close()
        cached.close()
    return blocks, {"kind": "worksheets", "entries": entries}


def parse_powerpoint(content: bytes) -> tuple[list[dict], dict]:
    """Inventory all slides and extract text/tables/notes without claiming visual coverage."""

    from pptx import Presentation

    presentation = Presentation(io.BytesIO(content))
    blocks = []
    entries = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        first_block = len(blocks)
        for shape_number, shape in enumerate(slide.shapes, 1):
            position = {"slide_number": slide_number, "shape_number": shape_number}
            if shape.has_text_frame:
                append_block(blocks, shape.text, **position)
            elif shape.has_table:
                for row_number, row in enumerate(shape.table.rows, 1):
                    cells = [cell.text for cell in row.cells]
                    append_block(
                        blocks,
                        json.dumps(cells, ensure_ascii=False),
                        kind="table_row",
                        cells=cells,
                        row_number=row_number,
                        **position,
                    )
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                append_block(blocks, frame.text, slide_number=slide_number, is_notes=True)
        entries.append({"number": slide_number, "block_count": len(blocks) - first_block})
    return blocks, {"kind": "slides", "entries": entries}


def parse_request(request: dict) -> dict:
    """Validate the fixed protocol; keep old Office responses and dispatch PDF without commands."""

    if not isinstance(request, dict) or set(request) != {
        "protocol",
        "media_type",
        "content_base64",
    }:
        raise ValueError("Invalid Office parse request")
    protocol = request["protocol"]
    if protocol not in ("office_parse.v1", "file_parse.v1"):
        raise ValueError("Unsupported file parse protocol")
    if not isinstance(request["media_type"], str) or not isinstance(request["content_base64"], str):
        raise ValueError("Invalid file request fields")
    content = base64.b64decode(request["content_base64"], validate=True)
    media_type = request["media_type"]
    if protocol == "file_parse.v1" and media_type == "application/pdf":
        # Isolated mode excludes the script directory from sys.path. Load only the
        # adjacent administrator-built module, never a caller-supplied module/path.
        import runpy

        from pathlib import Path

        parser = runpy.run_path(str(Path(__file__).with_name("parse_pdf.py")))
        pages, version = parser["extract_pages"](content)
        return {
            "protocol": protocol, "source_hash": hashlib.sha256(content).hexdigest(),
            "media_type": media_type, "parser_version": version, "pages": pages,
        }
    validate_package(content, media_type)
    parser = {
        "word/document.xml": parse_word,
        "xl/workbook.xml": parse_excel,
        "ppt/presentation.xml": parse_powerpoint,
    }[OFFICE_TYPES[media_type]]
    blocks, structure = parser(content)
    if not blocks:
        raise ValueError("Office document contains no supported text or table content")
    warnings = {
        parse_word: (
            "Body only: headers, footers, notes, revisions, nested tables, text boxes "
            "and images are not fully extracted; merged table cells may repeat text"
        ),
        parse_excel: (
            "Stored values only: formulas are not evaluated; charts, images, comments "
            "and hidden columns are not interpreted; empty sheets have no blocks"
        ),
        parse_powerpoint: (
            "Shape order is not visual reading order; groups, charts, images "
            "and merged table geometry are not interpreted"
        ),
    }
    versions = "/".join(
        f"{name}:{importlib.metadata.version(name)}"
        for name in ("python-docx", "openpyxl", "python-pptx")
    )
    return {
        "protocol": protocol,
        "source_hash": hashlib.sha256(content).hexdigest(),
        "media_type": media_type,
        "parser_version": f"office-v2/{versions}",
        "blocks": blocks,
        "status": "partial",
        "warnings": [warnings[parser]],
        **({"structure": structure} if structure is not None else {}),
    }


def limit_process_resources() -> None:
    """Allow private source files up to the upload limit; output is bounded separately."""
    resource.setrlimit(resource.RLIMIT_CPU, (30, 30))
    maximum = max(MAX_SOURCE_BYTES, MAX_RESULT_BYTES)
    resource.setrlimit(resource.RLIMIT_FSIZE, (maximum, maximum))


def main() -> int:
    """Bound source-file resources and result bytes, preserving safe Office error codes."""

    limit_process_resources()
    protocol = None
    try:
        encoded = sys.stdin.buffer.readline(MAX_REQUEST_BYTES + 1)
        if len(encoded) > MAX_REQUEST_BYTES or not encoded.endswith(b"\n"):
            raise ValueError("Office request size or framing is invalid")
        request = json.loads(encoded)
        protocol = request.get("protocol") if isinstance(request, dict) else None
        report = parse_request(request)
        output = json.dumps(report, ensure_ascii=False, allow_nan=False).encode()
        if len(output) > MAX_RESULT_BYTES:
            raise ValueError("Office result exceeds the supported size")
        sys.stdout.buffer.write(output + b"\n")
        return 0
    except Exception:
        # Library exceptions can include source text; never send their messages upstream.
        print(json.dumps({"error": "office_parse_failed" if protocol == "office_parse.v1"
                          else "file_parse_failed"}))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
