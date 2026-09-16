"""Measure real disposable Office-container latency and synthetic extraction coverage.

Run ``python3 deployments/file-runtime/benchmark_office.py --image IMAGE --repeat 3``
on macOS Docker Desktop or Linux Docker Engine. Uses only cached dependencies,
synthetic documents and isolated containers; does not call models or application DBs.
Times include create/start/parse/remove, not MinIO, upload transport or model generation.
"""

import argparse
import base64
import io
import json
import statistics
import subprocess
import sys
import time
import uuid
from pathlib import Path


def fixtures():
    """Generate text/table, image, header, chart and merged-cell evidence in a private container."""
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from PIL import Image
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    result = []

    def add(label, document, media_type, present, missing):
        """Serialize a synthetic document with independently declared expected markers."""
        target = io.BytesIO()
        document.save(target)
        result.append(
            {
                "label": label,
                "media_type": media_type,
                "content_base64": base64.b64encode(target.getvalue()).decode(),
                "present": present,
                "missing": missing,
            }
        )

    image = io.BytesIO()
    Image.new("RGB", (160, 80), "white").save(image, format="PNG")
    word = Document()
    word.add_heading("BODY_MARKER", 1)
    word.add_paragraph("中文正文与基础表格")
    word.add_table(rows=1, cols=2).cell(0, 0).text = "TABLE_MARKER"
    word.sections[0].header.paragraphs[0].text = "HEADER_MARKER"
    word.add_picture(io.BytesIO(image.getvalue()))
    add(
        "word_text_table_header_image",
        word,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ["BODY_MARKER", "TABLE_MARKER"],
        ["HEADER_MARKER"],
    )

    book = Workbook()
    sheet = book.active
    sheet.append(["MERGED_MARKER", 1])
    sheet.append(["ROW_MARKER", 2])
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "MERGE_CONTENT"
    sheet["C1"] = "=SUM(B1:B2)"
    chart = BarChart()
    chart.title = "CHART_MARKER"
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2))
    sheet.add_chart(chart, "E1")
    book.create_sheet("empty-hidden").sheet_state = "hidden"
    add(
        "excel_values_formula_merge_chart",
        book,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ["ROW_MARKER", "MERGE_CONTENT"],
        ["CHART_MARKER"],
    )

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(3), Inches(1)
    ).text = "SLIDE_MARKER"
    slide.notes_slide.notes_text_frame.text = "NOTES_MARKER"
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(
        io.BytesIO(image.getvalue()), Inches(1), Inches(1)
    )
    add("slides_text_notes_image", deck,
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ["SLIDE_MARKER", "NOTES_MARKER"], [])
    chart_data = ChartData()
    chart_data.categories = ["A", "B"]
    chart_data.add_series("CHART_MARKER", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(2),
        Inches(3),
        Inches(2),
        chart_data,
    )
    add(
        "slides_text_notes_chart_image",
        deck,
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ["SLIDE_MARKER", "NOTES_MARKER"],
        ["CHART_MARKER"],
    )
    return result


def benchmark(image, repeat):
    """Run each synthetic source through the fixed executor and report observed wall-clock times."""
    from manage_runtime import container_create_arguments
    from serve_runtime import parse_in_container, resolve_image

    image_id = resolve_image(image)
    name = f"nexuspilot-file-benchmark-{uuid.uuid4().hex}"
    args = container_create_arguments(name, image_id)
    args[-1:] = [
        "--interactive",
        "--entrypoint",
        "python",
        image_id,
        "-I",
        "-",
        "--fixtures",
    ]
    try:
        subprocess.run(args, check=True, capture_output=True, timeout=30)
        generated = subprocess.run(
            ["docker", "start", "--attach", "--interactive", name],
            input=Path(__file__).read_bytes(),
            capture_output=True,
            timeout=60,
            check=True,
        )
        sources = json.loads(generated.stdout)
    finally:
        subprocess.run(
            ["docker", "rm", "--force", name],
            check=True,
            capture_output=True,
            timeout=15,
        )
    reports = []
    for source in sources:
        request = (
            json.dumps(
                {
                    "protocol": "office_parse.v1",
                    "media_type": source["media_type"],
                    "content_base64": source["content_base64"],
                }
            ).encode()
            + b"\n"
        )
        timings = []
        for _ in range(repeat):
            started = time.perf_counter()
            report = parse_in_container(request, image_id)
            timings.append(round((time.perf_counter() - started) * 1000))
            if "blocks" not in report:
                continue
            text = "\n".join(block["text"] for block in report["blocks"])
            assert all(marker in text for marker in source["present"])
            assert all(marker not in text for marker in source["missing"])
        reports.append(
            {
                "sample": source["label"],
                "bytes": len(base64.b64decode(source["content_base64"])),
                "elapsed_ms": timings,
                "median_ms": statistics.median(timings),
                "blocks": len(report.get("blocks", [])),
                "status": report.get("status", "failed"),
                "error": report.get("error"),
                "extracted": source["present"] if "blocks" in report else [],
                "not_extracted": source["missing"] if "blocks" in report else source["present"] + source["missing"],
                "structure": report.get("structure"),
            }
        )
    print(
        json.dumps(
            {"image": image_id, "samples": reports}, ensure_ascii=False, indent=2
        )
    )


if __name__ == "__main__":
    if "--fixtures" in sys.argv:
        print(json.dumps(fixtures()))
    else:
        parser = argparse.ArgumentParser(description=__doc__)
        parser.add_argument("--image", default="nexuspilot-file-runtime:local")
        parser.add_argument("--repeat", type=int, choices=range(1, 11), default=3)
        options = parser.parse_args()
        benchmark(options.image, options.repeat)
